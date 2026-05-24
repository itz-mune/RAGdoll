"""RAGdoll sidecar — FastAPI server on localhost:8765."""
# ── Silence noisy third-party warnings before any HF imports ─────────────────
import os, warnings
os.environ.setdefault("HF_HUB_DISABLE_IMPLICIT_TOKEN", "1")   # no "unauthenticated" nag
os.environ.setdefault("TOKENIZERS_PARALLELISM", "false")       # avoids fork deadlock warning
warnings.filterwarnings("ignore", message=".*unauthenticated.*")
warnings.filterwarnings("ignore", message=".*HF_TOKEN.*")
# ─────────────────────────────────────────────────────────────────────────────
from contextlib import asynccontextmanager
import asyncio
import csv
import hashlib
import io
import json
import time
import uuid
from pathlib import Path
from typing import Optional, List

import httpx
import uvicorn
from fastapi import BackgroundTasks, Depends, FastAPI, HTTPException, Query, UploadFile, File
from fastapi.middleware.cors import CORSMiddleware
from fastapi.middleware.gzip import GZipMiddleware
from fastapi.responses import FileResponse, StreamingResponse
from pydantic import BaseModel
from sqlmodel import Session, select

from db import ConversationModel, ConversationSummary, MessageModel, engine, init_db
from permissions_bus import permission_context, respond_to_permission

# ── Global state ──────────────────────────────────────────────────────────────

_sidecar_stage: str = "starting"
_http_client: httpx.AsyncClient | None = None

# Pre-fetch cache: md5(query) → RetrievalResult (short-lived, used within same request window)
_prefetch_cache: dict[str, object] = {}


# ── Lifespan ──────────────────────────────────────────────────────────────────

@asynccontextmanager
async def lifespan(app: FastAPI):
    global _sidecar_stage, _http_client
    print("[RAGdoll] Sidecar starting on port 8765")

    _sidecar_stage = "db"
    init_db()
    print("[RAGdoll] Database initialised")

    # Shared HTTP/2 client pool
    _http_client = httpx.AsyncClient(
        http2=True,
        limits=httpx.Limits(max_connections=20, max_keepalive_connections=10),
        timeout=30.0,
    )

    # Kick off background tasks — don't block the port from opening
    asyncio.create_task(_cleanup_temp_docs())
    _sidecar_stage = "warming_up"
    asyncio.create_task(_warm_embed_model_bg())
    asyncio.create_task(_init_file_index_bg())

    # Yield immediately so uvicorn opens the port and /health starts responding
    yield

    print("[RAGdoll] Sidecar shutting down")
    if _http_client and not _http_client.is_closed:
        await _http_client.aclose()


def _get_http() -> httpx.AsyncClient:
    if _http_client is None or _http_client.is_closed:
        return httpx.AsyncClient(timeout=30.0)
    return _http_client


async def _cleanup_temp_docs():
    try:
        from memory.store import cleanup_temp_documents
        await asyncio.to_thread(cleanup_temp_documents)
    except Exception as exc:
        print(f"[RAGdoll] Temp doc cleanup failed: {exc}")


async def _warm_embed_model():
    try:
        from memory.embedder import warmup
        await asyncio.to_thread(warmup)
        print("[RAGdoll] Embedding model ready")
    except Exception as exc:
        print(f"[RAGdoll] Embedding model load failed: {exc}")


async def _warm_embed_model_bg():
    """Run model warm-up as a background task so uvicorn opens the port immediately."""
    global _sidecar_stage
    await _warm_embed_model()
    _sidecar_stage = "ready"


async def _init_file_index_bg():
    """Initialise the Universal File Access index after startup (best-effort)."""
    try:
        from plugins.state import get_plugin_config
        cfg = get_plugin_config("universal-file-access")
        from plugins.loader import _plugins_dir
        plugin_dir = _plugins_dir() / "universal-file-access"
        if not plugin_dir.exists():
            return  # plugin not installed yet
        import importlib.util, sys as _sys
        if str(plugin_dir) not in _sys.path:
            _sys.path.insert(0, str(plugin_dir))
        spec = importlib.util.spec_from_file_location("_uf_indexer", plugin_dir / "indexer.py")
        if spec and spec.loader:
            mod = importlib.util.module_from_spec(spec)
            spec.loader.exec_module(mod)  # type: ignore[union-attr]
            await mod.init_index(cfg)
    except Exception as exc:
        print(f"[RAGdoll] File index init skipped: {exc}")


# ── App ───────────────────────────────────────────────────────────────────────

app = FastAPI(title="RAGdoll Sidecar", version="0.1.0", lifespan=lifespan)

app.add_middleware(GZipMiddleware, minimum_size=500)
app.add_middleware(
    CORSMiddleware,
    allow_origins=["http://localhost:1420", "tauri://localhost"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)


def get_db():
    with Session(engine) as session:
        yield session


# ── Pydantic schemas ──────────────────────────────────────────────────────────

class SettingsValidationRequest(BaseModel):
    provider: str
    api_key: str
    model: Optional[str] = None


class ConversationCreate(BaseModel):
    title: str
    provider: str
    model: str


class ConversationUpdate(BaseModel):
    title: str


class ChatStreamRequest(BaseModel):
    conversation_id: str
    message: str
    display_message: Optional[str] = None
    provider: str
    api_key: Optional[str] = None
    model: Optional[str] = None
    file_names: List[str] = []
    response_style: str = "normal"
    has_response_style: bool = False


class FileProcessRequest(BaseModel):
    store_permanently: bool = False


# ══════════════════════════════════════════════════════════════════════════════
# Health
# ══════════════════════════════════════════════════════════════════════════════

@app.get("/health")
async def health() -> dict:
    return {"status": "ok", "version": "0.1.0", "stage": _sidecar_stage}


# ══════════════════════════════════════════════════════════════════════════════
# Settings / API-key validation
# ══════════════════════════════════════════════════════════════════════════════

@app.post("/settings/validate")
async def validate_settings(request: SettingsValidationRequest) -> dict:
    provider = request.provider.lower()
    api_key = request.api_key
    model = request.model
    try:
        if provider == "openai":
            return await _validate_openai(api_key)
        elif provider == "anthropic":
            return await _validate_anthropic(api_key)
        elif provider == "groq":
            return await _validate_groq(api_key)
        elif provider == "google":
            return await _validate_google(api_key)
        elif provider == "huggingface":
            if not model:
                return {"valid": False, "error": "Model ID required for HuggingFace"}
            return await _validate_huggingface(api_key, model)
        elif provider == "openrouter":
            if not model:
                return {"valid": False, "error": "Model slug required for OpenRouter"}
            return await _validate_openrouter(api_key, model)
        elif provider == "ollama":
            return await _validate_ollama(api_key)
        else:
            return {"valid": False, "error": f"Unknown provider: {provider}"}
    except Exception as exc:
        return {"valid": False, "error": str(exc)}


async def _validate_openai(api_key: str) -> dict:
    try:
        r = await _get_http().get(
            "https://api.openai.com/v1/models",
            headers={"Authorization": f"Bearer {api_key}"},
            timeout=10,
        )
        if r.status_code == 200:
            return {"valid": True, "error": None}
        if r.status_code == 401:
            return {"valid": False, "error": "Invalid API key"}
        return {"valid": False, "error": f"OpenAI error {r.status_code}"}
    except Exception as exc:
        return {"valid": False, "error": f"Connection error: {exc}"}


async def _validate_anthropic(api_key: str) -> dict:
    try:
        r = await _get_http().get(
            "https://api.anthropic.com/v1/models",
            headers={"x-api-key": api_key, "anthropic-version": "2023-06-01"},
            timeout=10,
        )
        if r.status_code == 200:
            return {"valid": True, "error": None}
        if r.status_code == 401:
            return {"valid": False, "error": "Invalid API key"}
        return {"valid": False, "error": f"Anthropic error {r.status_code}"}
    except Exception as exc:
        return {"valid": False, "error": f"Connection error: {exc}"}


async def _validate_groq(api_key: str) -> dict:
    try:
        r = await _get_http().get(
            "https://api.groq.com/openai/v1/models",
            headers={"Authorization": f"Bearer {api_key}"},
            timeout=10,
        )
        if r.status_code == 200:
            return {"valid": True, "error": None}
        if r.status_code == 401:
            return {"valid": False, "error": "Invalid API key"}
        return {"valid": False, "error": f"Groq error {r.status_code}"}
    except Exception as exc:
        return {"valid": False, "error": f"Connection error: {exc}"}


async def _validate_google(api_key: str) -> dict:
    try:
        r = await _get_http().get(
            f"https://generativelanguage.googleapis.com/v1beta/models?key={api_key}",
            timeout=10,
        )
        if r.status_code == 200:
            return {"valid": True, "error": None}
        if r.status_code in (400, 401):
            return {"valid": False, "error": "Invalid API key"}
        return {"valid": False, "error": f"Google error {r.status_code}"}
    except Exception as exc:
        return {"valid": False, "error": f"Connection error: {exc}"}


async def _validate_huggingface(api_key: str, model: str) -> dict:
    try:
        r = await _get_http().get(
            f"https://api-inference.huggingface.co/models/{model}",
            headers={"Authorization": f"Bearer {api_key}"},
            timeout=10,
        )
        if r.status_code == 200:
            return {"valid": True, "error": None}
        if r.status_code == 401:
            return {"valid": False, "error": "Invalid HuggingFace token"}
        if r.status_code == 404:
            return {"valid": False, "error": "Model not found"}
        if r.status_code == 403:
            return {"valid": False, "error": "Model requires a Pro subscription"}
        return {"valid": False, "error": f"HuggingFace error {r.status_code}"}
    except Exception as exc:
        return {"valid": False, "error": f"Connection error: {exc}"}


async def _validate_openrouter(api_key: str, model: str) -> dict:
    try:
        r = await _get_http().get(
            "https://openrouter.ai/api/v1/models",
            headers={"Authorization": f"Bearer {api_key}"},
            timeout=10,
        )
        if r.status_code == 401:
            return {"valid": False, "error": "Invalid OpenRouter key"}
        if r.status_code != 200:
            return {"valid": False, "error": f"OpenRouter error {r.status_code}"}
        model_ids = [m.get("id") for m in r.json().get("data", [])]
        if model not in model_ids:
            return {"valid": False, "error": "Model not found on OpenRouter"}
        return {"valid": True, "error": None}
    except Exception as exc:
        return {"valid": False, "error": f"Connection error: {exc}"}


async def _validate_ollama(base_url: str) -> dict:
    try:
        r = await _get_http().get(f"{base_url}/api/tags", timeout=10)
        if r.status_code == 200:
            return {"valid": True, "error": None}
        return {"valid": False, "error": f"Ollama error {r.status_code}"}
    except Exception as exc:
        return {"valid": False, "error": f"Connection error: {exc}"}


# ══════════════════════════════════════════════════════════════════════════════
# File processing
# ══════════════════════════════════════════════════════════════════════════════

@app.post("/files/process")
async def process_files(
    files: List[UploadFile] = File(...),
    store_permanently: bool = Query(default=False),
) -> dict:
    results = []
    for upload in files:
        filename = upload.filename or "unknown"
        ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
        raw_bytes = await upload.read()
        content: str = ""
        error: Optional[str] = None
        chunks_stored: int = 0

        try:
            if ext in ("txt", "md"):
                content = raw_bytes.decode("utf-8", errors="replace")
            elif ext == "pdf":
                try:
                    from pypdf import PdfReader
                    reader = PdfReader(io.BytesIO(raw_bytes))
                    content = "\n\n".join(page.extract_text() or "" for page in reader.pages)
                except ImportError:
                    error = "pypdf not installed — run: uv add pypdf"
            elif ext == "docx":
                try:
                    from docx import Document
                    doc = Document(io.BytesIO(raw_bytes))
                    content = "\n".join(para.text for para in doc.paragraphs)
                except ImportError:
                    error = "python-docx not installed — run: uv add python-docx"
            elif ext == "csv":
                text = raw_bytes.decode("utf-8", errors="replace")
                reader_csv = csv.reader(io.StringIO(text))
                rows = list(reader_csv)
                if rows:
                    header = " | ".join(rows[0])
                    sep = " | ".join(["---"] * len(rows[0]))
                    body = "\n".join(" | ".join(row) for row in rows[1:])
                    content = f"{header}\n{sep}\n{body}"
            elif ext == "json":
                decoded = raw_bytes.decode("utf-8", errors="replace")
                parsed = json.loads(decoded)
                content = json.dumps(parsed, indent=2, ensure_ascii=False)
            else:
                error = f"Unsupported file type: .{ext}"

            if content and not error:
                # Always save raw bytes to disk so the document viewer can render the file.
                from memory.store import _db_path
                files_dir = Path(_db_path()) / "files"
                files_dir.mkdir(parents=True, exist_ok=True)
                (files_dir / filename).write_bytes(raw_bytes)

                # Always index chunks in LanceDB so RAG retrieval and source highlights work.
                # Chat attachments are marked is_temporary=True — they are excluded from the
                # Memory Browser document list and deleted on next sidecar startup.
                file_hash = hashlib.sha256(raw_bytes).hexdigest()
                from memory.store import write_document
                chunks_stored = await asyncio.to_thread(
                    write_document, filename, content, file_hash,
                    not store_permanently,  # is_temporary
                )

        except Exception as exc:
            error = str(exc)

        results.append({
            "filename": filename,
            "content": content,
            "char_count": len(content),
            "error": error,
            "chunks_stored": chunks_stored,
        })

    return {"files": results}


# ══════════════════════════════════════════════════════════════════════════════
# Conversations
# ══════════════════════════════════════════════════════════════════════════════

@app.post("/conversations")
async def create_conversation(
    body: ConversationCreate,
    db: Session = Depends(get_db),
) -> dict:
    now = int(time.time() * 1000)
    conv = ConversationModel(
        id=str(uuid.uuid4()),
        title=body.title,
        provider=body.provider,
        model=body.model,
        created_at=now,
        updated_at=now,
        message_count=0,
    )
    db.add(conv)
    db.commit()
    db.refresh(conv)
    return _conv_dict(conv)


@app.get("/conversations")
async def list_conversations(db: Session = Depends(get_db)) -> list:
    convs = db.exec(
        select(ConversationModel).order_by(ConversationModel.updated_at.desc())
    ).all()
    return [_conv_dict(c) for c in convs]


@app.delete("/conversations/{conversation_id}")
async def delete_conversation(
    conversation_id: str,
    db: Session = Depends(get_db),
) -> dict:
    conv = db.get(ConversationModel, conversation_id)
    if not conv:
        return {"error": "Not found"}
    msgs = db.exec(
        select(MessageModel).where(MessageModel.conversation_id == conversation_id)
    ).all()
    for m in msgs:
        db.delete(m)
    db.delete(conv)
    db.commit()
    # Also remove vector memory for this conversation
    try:
        from memory.store import delete_conversation_memory
        await asyncio.to_thread(delete_conversation_memory, conversation_id)
    except Exception:
        pass
    return {"status": "deleted"}


@app.get("/conversations/{conversation_id}/messages")
async def get_messages(
    conversation_id: str,
    db: Session = Depends(get_db),
) -> list:
    msgs = db.exec(
        select(MessageModel)
        .where(MessageModel.conversation_id == conversation_id)
        .order_by(MessageModel.created_at)
    ).all()
    return [_msg_dict(m) for m in msgs]


@app.post("/conversations/{conversation_id}/title")
async def update_title(
    conversation_id: str,
    body: ConversationUpdate,
    db: Session = Depends(get_db),
) -> dict:
    conv = db.get(ConversationModel, conversation_id)
    if not conv:
        return {"error": "Not found"}
    conv.title = body.title
    conv.updated_at = int(time.time() * 1000)
    db.add(conv)
    db.commit()
    return {"status": "updated"}


# ══════════════════════════════════════════════════════════════════════════════
# Memory endpoints
# ══════════════════════════════════════════════════════════════════════════════

@app.get("/memory/stats")
async def memory_stats() -> dict:
    from memory.store import get_memory_stats
    return await asyncio.to_thread(get_memory_stats)


@app.get("/memory/search")
async def memory_search(
    q: str = Query(...),
    limit: int = Query(default=20, le=50),
) -> list:
    from memory.store import search_all_memory
    return await asyncio.to_thread(search_all_memory, q, limit)


@app.get("/memory/documents")
async def list_documents() -> list:
    from memory.store import get_document_list
    return await asyncio.to_thread(get_document_list)


@app.delete("/memory/chunks/{chunk_id}")
async def delete_memory_chunk(chunk_id: str) -> dict:
    from memory.store import delete_memory
    await asyncio.to_thread(delete_memory, chunk_id)
    return {"status": "deleted", "id": chunk_id}


@app.delete("/memory/documents/{filename}")
async def delete_document_endpoint(filename: str) -> dict:
    from memory.store import delete_document
    await asyncio.to_thread(delete_document, filename)
    return {"status": "deleted", "filename": filename}


@app.get("/memory/documents/{filename}/text")
async def get_document_text_endpoint(filename: str) -> dict:
    from memory.store import get_document_text, _db_path

    # Try LanceDB first (permanently stored docs have chunks there)
    result = await asyncio.to_thread(get_document_text, filename)
    if result:
        return result

    # Fall back to raw bytes saved during any /files/process call
    file_path = Path(_db_path()) / "files" / filename
    if not file_path.exists():
        raise HTTPException(status_code=404, detail="Document not found on disk")

    try:
        raw_bytes = file_path.read_bytes()
        ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
        content = ""

        if ext in ("txt", "md"):
            content = raw_bytes.decode("utf-8", errors="replace")
        elif ext == "pdf":
            from pypdf import PdfReader
            reader = PdfReader(io.BytesIO(raw_bytes))
            content = "\n\n".join(page.extract_text() or "" for page in reader.pages)
        elif ext == "docx":
            from docx import Document as DocxDocument
            doc = DocxDocument(io.BytesIO(raw_bytes))
            content = "\n".join(para.text for para in doc.paragraphs)
        elif ext == "csv":
            text = raw_bytes.decode("utf-8", errors="replace")
            reader_csv = csv.reader(io.StringIO(text))
            rows = list(reader_csv)
            if rows:
                header = " | ".join(rows[0])
                sep = " | ".join(["---"] * len(rows[0]))
                body = "\n".join(" | ".join(row) for row in rows[1:])
                content = f"{header}\n{sep}\n{body}"
        elif ext == "json":
            decoded = raw_bytes.decode("utf-8", errors="replace")
            content = json.dumps(json.loads(decoded), indent=2, ensure_ascii=False)

        if not content:
            raise HTTPException(status_code=422, detail=f"Could not extract text from .{ext} file")

        return {"filename": filename, "text": content, "chunk_count": 0}
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(status_code=500, detail=f"Text extraction failed: {exc}")


@app.get("/memory/documents/{filename}/file")
async def get_document_file(filename: str) -> FileResponse:
    from memory.store import _db_path
    file_path = Path(_db_path()) / "files" / filename
    if not file_path.exists():
        raise HTTPException(status_code=404, detail="File not found")
    return FileResponse(str(file_path), filename=filename)


@app.get("/memory/recent")
async def memory_recent(limit: int = Query(default=20, le=50)) -> list:
    from memory.store import get_recent_memories
    return await asyncio.to_thread(get_recent_memories, limit)


@app.get("/memory/conversations")
async def memory_conversations() -> list:
    from memory.store import get_memory_by_conversation
    return await asyncio.to_thread(get_memory_by_conversation)


@app.post("/memory/compact")
async def compact_memory(background_tasks: BackgroundTasks) -> dict:
    from memory.compactor import run_full_compaction
    stats = await run_full_compaction()
    return {"status": "ok", "stats": stats}


@app.delete("/memory/all")
async def clear_all_memory() -> dict:
    from memory.store import clear_all_semantic_memory
    await asyncio.to_thread(clear_all_semantic_memory)
    return {"status": "cleared"}


@app.get("/memory/prefetch")
async def prefetch_context(q: str = Query(...)) -> dict:
    """Pre-warm retrieval for a likely query while the user is still typing."""
    import hashlib as _hl
    key = _hl.md5(q.encode()).hexdigest()
    try:
        from memory.retriever import retrieve_context
        result = await retrieve_context(q, conversation_id="")
        _prefetch_cache[key] = result
        # Evict oldest entries when cache grows large
        if len(_prefetch_cache) > 64:
            oldest = next(iter(_prefetch_cache))
            _prefetch_cache.pop(oldest, None)
    except Exception:
        pass
    return {"status": "ok", "key": key}


# ══════════════════════════════════════════════════════════════════════════════
# Debug / benchmark
# ══════════════════════════════════════════════════════════════════════════════

@app.get("/debug/benchmark")
async def benchmark() -> dict:
    """Benchmark embedding throughput and LRU cache hit rate."""
    import time as _t
    from memory.embedder import embed, lru_stats

    queries = [
        "What is machine learning?",
        "Tell me about transformer models",
        "Hello world test sentence",
        "RAGdoll is a local AI assistant",
    ]
    # Cold run (first call populates LRU cache)
    t0 = _t.perf_counter()
    for _ in range(5):
        await asyncio.to_thread(embed, queries)
    cold_ms = round((_t.perf_counter() - t0) * 1000, 1)

    # Hot run (all from LRU cache)
    t1 = _t.perf_counter()
    for _ in range(5):
        await asyncio.to_thread(embed, queries)
    hot_ms = round((_t.perf_counter() - t1) * 1000, 1)

    from cache.response_cache import stats as cache_stats
    return {
        "embed_cold_5runs_ms": cold_ms,
        "embed_hot_5runs_ms": hot_ms,
        "speedup_x": round(cold_ms / hot_ms, 1) if hot_ms > 0 else 0,
        "lru": lru_stats(),
        "response_cache": cache_stats(),
        "sidecar_stage": _sidecar_stage,
    }


# ══════════════════════════════════════════════════════════════════════════════
# Plugin endpoints
# ══════════════════════════════════════════════════════════════════════════════

class PluginInstallRequest(BaseModel):
    id: str
    files: list[dict]  # [{ path: str, content: str }]


class PluginConfigRequest(BaseModel):
    config: dict


@app.get("/plugins")
async def list_plugins() -> dict:
    from plugins.loader import get_all_plugins
    from plugins.state import get_plugin_state
    plugins = await asyncio.to_thread(get_all_plugins)
    state = await asyncio.to_thread(get_plugin_state)
    return {
        "plugins": [p.to_dict() for p in plugins],
        "active_style": state.active_style,
    }


@app.post("/plugins/{plugin_id}/enable")
async def enable_plugin_endpoint(plugin_id: str) -> dict:
    from plugins.loader import enable_plugin
    await asyncio.to_thread(enable_plugin, plugin_id)
    return {"status": "enabled", "id": plugin_id}


@app.post("/plugins/{plugin_id}/disable")
async def disable_plugin_endpoint(plugin_id: str) -> dict:
    from plugins.loader import disable_plugin
    await asyncio.to_thread(disable_plugin, plugin_id)
    return {"status": "disabled", "id": plugin_id}


@app.post("/plugins/{plugin_id}/uninstall")
async def uninstall_plugin_endpoint(plugin_id: str) -> dict:
    from plugins.loader import uninstall_plugin
    await asyncio.to_thread(uninstall_plugin, plugin_id)
    return {"status": "uninstalled", "id": plugin_id}


@app.get("/plugins/{plugin_id}/config")
async def get_plugin_config_endpoint(plugin_id: str) -> dict:
    from plugins.state import get_plugin_config
    config = await asyncio.to_thread(get_plugin_config, plugin_id)
    return {"config": config}


@app.post("/plugins/{plugin_id}/config")
async def set_plugin_config_endpoint(plugin_id: str, body: PluginConfigRequest) -> dict:
    from plugins.state import set_plugin_config
    await asyncio.to_thread(set_plugin_config, plugin_id, body.config)
    return {"success": True}


@app.post("/plugins/{plugin_id}/config/reset")
async def reset_plugin_config_endpoint(plugin_id: str) -> dict:
    from plugins.state import reset_plugin_config
    await asyncio.to_thread(reset_plugin_config, plugin_id)
    return {"success": True}


@app.post("/plugins/style/{plugin_id}/activate")
async def activate_style(plugin_id: str) -> dict:
    from plugins.state import set_active_style
    await asyncio.to_thread(set_active_style, plugin_id)
    return {"status": "activated", "id": plugin_id}


@app.post("/plugins/style/clear")
async def clear_style() -> dict:
    from plugins.state import set_active_style
    await asyncio.to_thread(set_active_style, None)
    return {"status": "cleared"}


@app.post("/plugins/install")
async def install_plugin_endpoint(body: PluginInstallRequest) -> dict:
    try:
        from plugins.loader import install_plugin
        await asyncio.to_thread(install_plugin, body.id, body.files)
        return {"success": True, "error": None}
    except Exception as exc:
        return {"success": False, "error": str(exc)}


# ══════════════════════════════════════════════════════════════════════════════
# Permission bus
# ══════════════════════════════════════════════════════════════════════════════

class PermissionResponseRequest(BaseModel):
    request_id: str
    approved: bool


@app.post("/permissions/respond")
async def permissions_respond(body: PermissionResponseRequest) -> dict:
    found = respond_to_permission(body.request_id, body.approved)
    return {"ok": found}


# ══════════════════════════════════════════════════════════════════════════════
# Universal File Access index management
# ══════════════════════════════════════════════════════════════════════════════

@app.get("/plugins/universal-file-access/index/stats")
async def uf_index_stats() -> dict:
    try:
        from plugins.loader import _plugins_dir
        import sys as _sys, importlib.util as _ilu
        plugin_dir = _plugins_dir() / "universal-file-access"
        if str(plugin_dir) not in _sys.path:
            _sys.path.insert(0, str(plugin_dir))
        spec = _ilu.spec_from_file_location("_uf_indexer_s", plugin_dir / "indexer.py")
        if not spec or not spec.loader:
            return {"error": "plugin not installed"}
        mod = _ilu.module_from_spec(spec)
        spec.loader.exec_module(mod)  # type: ignore[union-attr]
        return mod.get_index_stats()
    except Exception as exc:
        raise HTTPException(status_code=500, detail=str(exc))


@app.post("/plugins/universal-file-access/index/rebuild")
async def uf_index_rebuild(background_tasks: BackgroundTasks) -> dict:
    async def _do_rebuild():
        try:
            from plugins.loader import _plugins_dir
            from plugins.state import get_plugin_config
            import sys as _sys, importlib.util as _ilu
            plugin_dir = _plugins_dir() / "universal-file-access"
            if str(plugin_dir) not in _sys.path:
                _sys.path.insert(0, str(plugin_dir))
            spec = _ilu.spec_from_file_location("_uf_indexer_r", plugin_dir / "indexer.py")
            if spec and spec.loader:
                mod = _ilu.module_from_spec(spec)
                spec.loader.exec_module(mod)  # type: ignore[union-attr]
                cfg = get_plugin_config("universal-file-access")
                stats = await mod.build_index(cfg)
                print(f"[RAGdoll] File index rebuilt: {stats.files_indexed} files")
        except Exception as exc:
            print(f"[RAGdoll] Index rebuild failed: {exc}")

    asyncio.create_task(_do_rebuild())
    return {"status": "building"}


@app.delete("/plugins/universal-file-access/index")
async def uf_index_clear() -> dict:
    try:
        from plugins.loader import _plugins_dir
        import sys as _sys, importlib.util as _ilu
        plugin_dir = _plugins_dir() / "universal-file-access"
        if str(plugin_dir) not in _sys.path:
            _sys.path.insert(0, str(plugin_dir))
        spec = _ilu.spec_from_file_location("_uf_indexer_c", plugin_dir / "indexer.py")
        if spec and spec.loader:
            mod = _ilu.module_from_spec(spec)
            spec.loader.exec_module(mod)  # type: ignore[union-attr]
            mod.clear_index()
        return {"ok": True}
    except Exception as exc:
        raise HTTPException(status_code=500, detail=str(exc))


# ══════════════════════════════════════════════════════════════════════════════
# Dashboard
# ══════════════════════════════════════════════════════════════════════════════

@app.get("/dashboard/stats")
async def dashboard_stats(days: int = 30) -> dict:
    try:
        from plugins.activity import get_stats
        return await asyncio.to_thread(get_stats, days)
    except Exception as exc:
        raise HTTPException(status_code=500, detail=str(exc))


@app.get("/dashboard/activity")
async def dashboard_activity(limit: int = 20) -> list:
    try:
        from plugins.activity import get_recent_activity
        return await asyncio.to_thread(get_recent_activity, limit)
    except Exception as exc:
        raise HTTPException(status_code=500, detail=str(exc))


class PinnedRequest(BaseModel):
    ids: List[str]

@app.post("/dashboard/pinned")
async def dashboard_pinned(body: PinnedRequest) -> list:
    try:
        from plugins.activity import get_pinned_conversations
        return await asyncio.to_thread(get_pinned_conversations, body.ids)
    except Exception as exc:
        raise HTTPException(status_code=500, detail=str(exc))


# ══════════════════════════════════════════════════════════════════════════════
# Chat streaming
# ══════════════════════════════════════════════════════════════════════════════

@app.post("/chat/stream")
async def chat_stream(request: ChatStreamRequest):
    async def generate():
        with Session(engine) as db:
            async for chunk in _stream_response(
                db=db,
                conversation_id=request.conversation_id,
                user_message=request.message,
                display_message=request.display_message,
                provider=request.provider,
                api_key=request.api_key or "",
                model=request.model,
                file_names=request.file_names,
                response_style=request.response_style,
                has_response_style=request.has_response_style,
            ):
                yield chunk

    return StreamingResponse(
        generate(),
        media_type="text/event-stream",
        headers={"Cache-Control": "no-cache", "X-Accel-Buffering": "no"},
    )


async def _store_memory(
    conversation_id: str,
    message_id: str,
    text: str,
    role: str,
    importance_override,
) -> None:
    try:
        from memory.store import write_memory
        ids = await asyncio.to_thread(
            write_memory, conversation_id, message_id, text, role, importance_override
        )
        print(f"[Memory] Stored {len(ids)} chunk(s) for {role} msg {message_id[:8]}")
    except Exception as exc:
        import traceback
        print(f"[Memory] write_memory failed ({role}): {exc}")
        traceback.print_exc()


async def _stream_response(
    db: Session,
    conversation_id: str,
    user_message: str,
    display_message: Optional[str],
    provider: str,
    api_key: str,
    model: Optional[str],
    file_names: List[str] = None,
    response_style: str = "normal",
    has_response_style: bool = False,
):
    from langchain_core.messages import AIMessage, HumanMessage, SystemMessage

    try:
        now = int(time.time() * 1000)
        conv = db.get(ConversationModel, conversation_id)
        if not conv:
            yield _sse({"type": "error", "message": f"Conversation {conversation_id!r} not found"})
            return

        # 1. History (last 10 turns before this message)
        history = db.exec(
            select(MessageModel)
            .where(MessageModel.conversation_id == conversation_id)
            .order_by(MessageModel.created_at.desc())
            .limit(10)
        ).all()
        history.reverse()

        # 2. Save user message
        user_msg_id = str(uuid.uuid4())
        db.add(MessageModel(
            id=user_msg_id,
            conversation_id=conversation_id,
            role="user",
            content=user_message,
            display_content=display_message if display_message is not None else user_message,
            attached_file_names=json.dumps(file_names or []) if file_names else None,
            created_at=now,
        ))
        db.commit()

        # 3. Async write user message to memory (fire-and-forget)
        from memory.selective_storage import storage_decision
        user_storage = storage_decision(user_message, "user")
        if user_storage["should_store"]:
            asyncio.create_task(_store_memory(
                conversation_id, user_msg_id, user_message, "user",
                user_storage["importance_override"],
            ))

        # 4. Route query + parallel retrieval / summary / plugin load
        from memory.router import route_query
        from memory.retriever import retrieve_context, RetrievalResult
        from plugins.loader import get_active_style, get_enabled_skills

        retrieval_query = display_message if display_message else user_message
        route = route_query(retrieval_query)

        # Check pre-fetch cache first
        import hashlib as _hl
        _pf_key = _hl.md5(retrieval_query.encode()).hexdigest()
        _prefetch_hit = _prefetch_cache.pop(_pf_key, None)

        async def _do_retrieval() -> RetrievalResult:
            if _prefetch_hit is not None:
                return _prefetch_hit  # type: ignore[return-value]
            if route == "skip":
                return RetrievalResult()
            coro = retrieve_context(retrieval_query, conversation_id)
            if route == "retrieve":
                try:
                    return await asyncio.wait_for(coro, timeout=1.5)
                except asyncio.TimeoutError:
                    return RetrievalResult()
            return await coro  # "force"

        retrieval, summary, plugin_tools, active_style = await asyncio.gather(
            _do_retrieval(),
            asyncio.to_thread(_get_conversation_summary, db, conversation_id),
            asyncio.to_thread(get_enabled_skills),
            asyncio.to_thread(get_active_style),
        )

        # 5. Contextual compression
        from memory.compressor import compress_chunks
        retrieval.chunks = compress_chunks(retrieval.chunks, retrieval_query, provider)

        # 6. Build system prompt
        from graph.agent import get_style_contract
        from memory.prompt_builder import build_system_prompt
        style_suffix = get_style_contract(response_style) if has_response_style else ""
        if active_style and active_style.get("system_prompt_prefix"):
            style_suffix = active_style["system_prompt_prefix"] + "\n\n" + style_suffix

        # Inject capability-expansion policy when skill-finder is active
        has_skill_finder = any(
            getattr(t, "name", "") in ("search_marketplace", "install_skill", "install_and_activate_style")
            for t in (plugin_tools or [])
        )
        if has_skill_finder:
            skill_finder_policy = (
                "CAPABILITY EXPANSION POLICY (applies to every message):\n"
                "You have access to a plugin marketplace via the search_marketplace tool.\n"
                "Whenever the user's request falls outside your native abilities — or when they ask for\n"
                "a speaking style (pirate, Shakespearean, formal, etc.), an image generator, a web search,\n"
                "a video transcript, file processing, or ANY specialised capability — follow this exact workflow:\n"
                "  1. Call search_marketplace(query) with a short description of what is needed.\n"
                "  2. If a matching plugin is found: call install_skill or install_and_activate_style with its id.\n"
                "  3. If no plugin matches: tell the user nothing was found, then do your best natively.\n"
                "NEVER skip step 1. NEVER mimic a style or simulate a capability without searching first.\n"
                "For styles specifically: calling install_and_activate_style makes the style persist across\n"
                "ALL future messages. Just responding in the style yourself only lasts one reply."
            )
            style_suffix = skill_finder_policy + ("\n\n" + style_suffix if style_suffix else "")

        # Inject file-search policy when universal-file-access is active.
        # Without this the LLM falls back on its safety training ("I don't have
        # access to your file system") even when the search_files tool is loaded.
        has_file_search = any(
            getattr(t, "name", "") == "search_files"
            for t in (plugin_tools or [])
        )
        if has_file_search:
            file_search_policy = (
                "FILE SEARCH POLICY (applies to every message):\n"
                "You have a search_files tool that can find any file on the user's computer.\n"
                "ALWAYS call search_files when the user asks you to find, locate, open, or access\n"
                "a file, document, image, video, spreadsheet, or any local item — even if they say\n"
                "things like 'on my desktop', 'in my downloads', 'the one I made last week', etc.\n"
                "NEVER say you cannot access the file system. NEVER ask the user to find the file\n"
                "themselves. NEVER explain how to use File Explorer or Finder.\n"
                "Instead: call search_files(query) immediately. The tool handles permissions."
            )
            style_suffix = file_search_policy + ("\n\n" + style_suffix if style_suffix else "")

        system_prompt_text = build_system_prompt(retrieval, summary, style_suffix)

        # 7. Check semantic response cache
        system_hash = hashlib.md5(system_prompt_text.encode()).hexdigest()
        from cache.response_cache import lookup as cache_lookup, store as cache_store
        cache_hit = cache_lookup(retrieval_query, system_hash)

        if cache_hit:
            # Replay cached response
            if cache_hit.thinking:
                yield _sse({"type": "thinking", "content": cache_hit.thinking, "is_partial": False})
                yield _sse({"type": "thinking_done"})
            if cache_hit.tool_calls:
                yield _sse({"type": "tool_use", "tools": cache_hit.tool_calls})
            for chunk in _chunk_for_display(cache_hit.response):
                yield _sse_chunk(chunk)
            assistant_msg_id = str(uuid.uuid4())
            db.add(MessageModel(
                id=assistant_msg_id,
                conversation_id=conversation_id,
                role="assistant",
                content=cache_hit.response,
                created_at=int(time.time() * 1000),
                tool_calls_made=json.dumps(cache_hit.tool_calls) if cache_hit.tool_calls else None,
            ))
            conv.message_count += 2
            conv.updated_at = int(time.time() * 1000)
            db.add(conv)
            db.commit()
            yield _sse({"type": "done", "messageId": assistant_msg_id})
            return

        # 8. Build LangChain message list — use cache_control blocks for Anthropic
        known_files: list = file_names or []
        if provider.lower() == "anthropic":
            sys_content = [{"type": "text", "text": system_prompt_text, "cache_control": {"type": "ephemeral"}}]
            lc_messages = [SystemMessage(content=sys_content)]
        else:
            lc_messages = [SystemMessage(content=system_prompt_text)]

        for m in history:
            if m.role == "user":
                lc_messages.append(HumanMessage(content=m.content))
            elif m.role == "assistant":
                lc_messages.append(AIMessage(content=m.content))
        lc_messages.append(HumanMessage(content=user_message))

        # 9. Emit memory context event
        if retrieval.chunks:
            yield _sse({
                "type": "memory",
                "chunks": [
                    {"id": c.id, "text": c.text[:1400], "score": round(c.score, 3), "source": c.source}
                    for c in retrieval.chunks
                ],
                "semantic": retrieval.semantic_count,
                "docs": retrieval.document_count,
            })

        # 10. Run the LangGraph — drain event_queue concurrently for early tool_use events
        from graph.agent import build_chat_graph, get_style_contract as _gsc
        event_queue: asyncio.Queue = asyncio.Queue()
        graph = build_chat_graph(
            provider, api_key, model, response_style, has_response_style,
            tools=plugin_tools if plugin_tools else None,
            event_queue=event_queue,
        )
        initial_state = {
            "messages": lc_messages,
            "has_style": has_response_style,
            "style": response_style,
            "style_contract": _gsc(response_style),
            "original_answer": "",
            "styled_answer": "",
            "final_answer": "",
            "thinking_content": "",
            "tool_calls_made": [],
            "style_score": 0.0,
            "style_feedback": "",
            "style_iteration": 0,
            "max_style_iterations": 5,
        }

        assistant_msg_id = str(uuid.uuid4())
        _graph_start_ms = int(time.time() * 1000)
        final_state_box: list = [initial_state]

        async def _run_graph():
            async for state in graph.astream(initial_state, stream_mode="values"):
                final_state_box[0] = state

        # Make the permission bus available inside the graph task.
        # asyncio.create_task() copies the current Context, so the ContextVar
        # is automatically inherited by the task (and by any tool _arun calls).
        permission_context.set(event_queue)
        graph_task = asyncio.create_task(_run_graph())
        while not graph_task.done():
            try:
                event = event_queue.get_nowait()
                yield _sse(event)
            except asyncio.QueueEmpty:
                await asyncio.sleep(0.02)
        await graph_task

        while not event_queue.empty():
            yield _sse(event_queue.get_nowait())

        final_state = final_state_box[0]
        full_content = final_state.get("final_answer", "")
        thinking_content = final_state.get("thinking_content", "")
        tool_calls_made: list = final_state.get("tool_calls_made", [])
        if not full_content:
            messages_list = final_state.get("messages", [])
            if messages_list:
                raw = getattr(messages_list[-1], "content", "")
                full_content = _message_content_to_text(raw)

        # 11. Emit thinking + stream chunks (no sleep — rAF batches on the client)
        if thinking_content:
            yield _sse({"type": "thinking", "content": thinking_content, "is_partial": False})
            yield _sse({"type": "thinking_done"})

        for content_chunk in _chunk_for_display(full_content):
            yield _sse_chunk(content_chunk)

        # 12. Citations
        citation_set: set[str] = set(known_files or [])
        for c in retrieval.chunks:
            if c.source != "conversation":
                citation_set.add(c.source)
        citations = sorted(citation_set)
        if citations:
            yield _sse({"type": "citations", "citations": citations})

        # 13. Persist assistant message + extract token usage
        _response_time_ms = int(time.time() * 1000) - _graph_start_ms
        _tokens_used: Optional[int] = None
        try:
            _last_msg = final_state.get("messages", [])[-1] if final_state.get("messages") else None
            if _last_msg and hasattr(_last_msg, "usage_metadata") and _last_msg.usage_metadata:
                _tokens_used = _last_msg.usage_metadata.get("total_tokens")
            elif _last_msg and hasattr(_last_msg, "response_metadata"):
                _rm = _last_msg.response_metadata or {}
                _tu = _rm.get("token_usage") or _rm.get("usage") or {}
                _tokens_used = _tu.get("total_tokens") or _tu.get("input_tokens", 0) + _tu.get("output_tokens", 0) or None
        except Exception:
            pass

        db.add(MessageModel(
            id=assistant_msg_id,
            conversation_id=conversation_id,
            role="assistant",
            content=full_content,
            created_at=int(time.time() * 1000),
            memory_chunks=json.dumps([
                {"id": c.id, "text": c.text[:1400], "score": c.score, "source": c.source}
                for c in retrieval.chunks
            ]) if retrieval.chunks else None,
            citations=json.dumps(citations) if citations else None,
            tool_calls_made=json.dumps(tool_calls_made) if tool_calls_made else None,
            tokens_used=_tokens_used,
            response_time_ms=_response_time_ms,
        ))
        conv.message_count += 2
        conv.updated_at = int(time.time() * 1000)
        db.add(conv)
        db.commit()

        # 14. Store response in semantic cache + write assistant memory (fire-and-forget)
        asyncio.create_task(asyncio.to_thread(
            cache_store, retrieval_query, system_hash, full_content, thinking_content, tool_calls_made
        ))
        asst_storage = storage_decision(full_content, "assistant")
        if asst_storage["should_store"]:
            asyncio.create_task(_store_memory(
                conversation_id, assistant_msg_id, full_content, "assistant",
                asst_storage["importance_override"],
            ))

        # 15. Auto-title on first exchange
        if conv.message_count == 2 and conv.title in ("New conversation", ""):
            auto_title = await _generate_title(provider, api_key, model, user_message)
            if auto_title:
                conv.title = auto_title
                db.add(conv)
                db.commit()
                yield _sse({"type": "title", "title": auto_title})

        # 16. Schedule compaction check
        from memory.compactor import record_message_time, maybe_compact
        record_message_time(conversation_id)
        asyncio.create_task(maybe_compact(conversation_id))

        yield _sse({"type": "done", "messageId": assistant_msg_id})

    except Exception as exc:
        print(f"[RAGdoll] Stream error: {exc}")
        yield _sse({"type": "error", "message": str(exc)})


# ── Helpers ───────────────────────────────────────────────────────────────────

def _sse(payload: dict) -> str:
    return f"data: {json.dumps(payload)}\n\n"


def _sse_chunk(text: str) -> str:
    """Compact SSE for text chunks — smaller wire payload, parsed by frontend."""
    return f'data: {{"t":"c","d":{json.dumps(text)}}}\n\n'


def _message_content_to_text(raw) -> str:
    if isinstance(raw, list):
        return "".join(
            block.get("text", "")
            for block in raw
            if isinstance(block, dict) and block.get("type") == "text"
        )
    return raw if isinstance(raw, str) else str(raw or "")


def _chunk_for_display(text: str, target_size: int = 28):
    if not text:
        return
    start = 0
    while start < len(text):
        end = min(start + target_size, len(text))
        if end < len(text):
            boundary = max(text.rfind(" ", start, end), text.rfind("\n", start, end))
            if boundary > start + 8:
                end = boundary + 1
        yield text[start:end]
        start = end


def _conv_dict(c: ConversationModel) -> dict:
    return {
        "id": c.id,
        "title": c.title,
        "provider": c.provider,
        "model": c.model,
        "createdAt": c.created_at,
        "updatedAt": c.updated_at,
        "messageCount": c.message_count,
    }


def _msg_dict(m: MessageModel) -> dict:
    return {
        "id": m.id,
        "conversationId": m.conversation_id,
        "role": m.role,
        "content": m.content,
        "displayContent": m.display_content,
        "attachedFileNames": json.loads(m.attached_file_names) if m.attached_file_names else None,
        "createdAt": m.created_at,
        "isStreaming": False,
        "memoryChunks": json.loads(m.memory_chunks) if m.memory_chunks else None,
        "citations": json.loads(m.citations) if m.citations else None,
        "toolCallsMade": json.loads(m.tool_calls_made) if m.tool_calls_made else None,
    }


def _get_conversation_summary(db: Session, conversation_id: str) -> Optional[str]:
    summary = db.exec(
        select(ConversationSummary)
        .where(ConversationSummary.conversation_id == conversation_id)
        .order_by(ConversationSummary.created_at.desc())
        .limit(1)
    ).first()
    return summary.summary_text if summary else None


async def _generate_title(
    provider: str, api_key: str, model: Optional[str], user_message: str
) -> str:
    try:
        from graph.agent import get_llm
        from langchain_core.messages import HumanMessage, SystemMessage
        llm = get_llm(provider, api_key, model)
        response = await llm.ainvoke([
            SystemMessage(content=(
                "Generate a concise 3-6 word title for a chat conversation "
                "based on the user's first message. Return only the title — "
                "no quotes, no period at the end."
            )),
            HumanMessage(content=user_message[:400]),
        ])
        raw = response.content
        if isinstance(raw, list):
            title = "".join(
                b.get("text", "") for b in raw
                if isinstance(b, dict) and b.get("type") == "text"
            )
        else:
            title = str(raw)
        return title.strip().strip('"').strip("'")[:60]
    except Exception as exc:
        print(f"[RAGdoll] Title generation failed: {exc}")
        return " ".join(user_message.strip().split()[:6])


# ── Entry point ───────────────────────────────────────────────────────────────

def _free_port(port: int) -> None:
    import platform
    import subprocess
    system = platform.system()
    try:
        if system == "Windows":
            result = subprocess.run(
                ["netstat", "-ano"],
                capture_output=True, text=True, timeout=5,
            )
            for line in result.stdout.splitlines():
                if f":{port}" in line and "LISTENING" in line:
                    pid = line.split()[-1]
                    subprocess.run(["taskkill", "/F", "/PID", pid], capture_output=True, timeout=5)
                    print(f"[RAGdoll] Freed port {port} (killed PID {pid})")
        else:
            result = subprocess.run(
                ["lsof", "-ti", f":{port}"],
                capture_output=True, text=True, timeout=5,
            )
            for pid in result.stdout.strip().splitlines():
                subprocess.run(["kill", "-9", pid], timeout=5)
                print(f"[RAGdoll] Freed port {port} (killed PID {pid})")
    except Exception as exc:
        print(f"[RAGdoll] Warning: could not free port {port}: {exc}")


if __name__ == "__main__":
    from config import get_sidecar_host, get_sidecar_port
    port = get_sidecar_port()
    host = get_sidecar_host()
    _free_port(port)

    run_kwargs: dict = {"host": host, "port": port, "log_level": "info"}

    # httptools is faster than the default h11 HTTP parser
    try:
        import httptools  # noqa: F401
        run_kwargs["http"] = "httptools"
    except ImportError:
        pass

    # uvloop is Linux/Mac only; Windows falls back to the default asyncio loop
    try:
        import uvloop  # noqa: F401
        import platform
        if platform.system() != "Windows":
            run_kwargs["loop"] = "uvloop"
    except ImportError:
        pass

    uvicorn.run(app, **run_kwargs)
